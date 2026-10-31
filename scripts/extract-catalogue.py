#!/usr/bin/env python3
"""Re-extract the product catalogue from VTech's sheet.

Source: Google Sheet "Ashutosh Project Status" (Drive id 1yWx6z6mkbqic78XTSxXg8Oe3wcD-CMYQ1J9dX6Omv08),
tab "Hansal". Download it as .xlsx (File → Download → Microsoft Excel), then:

    pip install openpyxl
    python3 scripts/extract-catalogue.py path/to/Ashutosh.xlsx
    node scripts/gen-catalogue.mjs          # → supabase/seed/catalogue.sql

Writes supabase/seed/products.csv and product_steps.csv. Colours (sheet legend, rows ~36-40):
green D9EAD3 + blue C9DAF8 = other agencies (external); dark red E06666 = VTech Installers;
light red F4CCCC = VTech Programmers; orange F9CB9C = VTech Project Managers. Green = phase 1; everything else = phase 2
(unlocks together once all greens of that product are Done/N.A.). Unknown colour → script stops so a human decides.
"""
import csv, re, sys
from pathlib import Path
import openpyxl

COL = {'FFD9EAD3': ('external', 1), 'FFC9DAF8': ('external', 2), 'FFCFE2F3': ('external', 2),
       'FFE06666': ('internal', 2), 'FFF4CCCC': ('internal', 2), 'FFF9CB9C': ('internal', 2)}
INTERNAL_TYPE = {'FFE06666': 'VTech Installers', 'FFF4CCCC': 'VTech Programmers', 'FFF9CB9C': 'VTech Project Managers'}
FIX = {'Fabric installtion': 'Fabric installation', 'Partial installtion done': 'Partial installation done',
       'Light Adddressing (Partial)': 'Light addressing (partial)', 'Light Adddressing ( complete )': 'Light addressing (complete)',
       'Readiness by carpentar': 'Readiness by carpenter', 'Ferrul Sheet': 'Ferrule sheet', 'Special prog.': 'Special programming'}
EXTRA_LOC = {'Wifi Access point': ['Ceiling', 'Room', 'Hall', 'Entrance', 'External']}

def ext_type(title, label):
    l = label.lower()
    if 'cable' in l or 'ferrule' in l: return 'Electrician'
    if 'dustfree' in l: return 'Painter'
    if 'carpenter' in l: return 'Carpenter'
    if 'opening details' in l or 'engraving' in l: return 'Architect'
    if 'schedule from the other vendor' in l: return 'AC Vendor' if 'AC' in title else 'Curtain Vendor'
    if 'fabric' in l: return 'Curtain Vendor'
    if 'internet' in l: return 'Network Agency'
    if l in ('partial installation done', 'complete installation'): return 'Electrician'
    if 'readiness to install' in l: return 'Electrician'
    return 'Client'  # delivery, client impressions, scene setting, AMC intro

def main(path):
    ws = openpyxl.load_workbook(path)['Hansal']
    prod, rows = [], []
    for r in range(3, ws.max_row + 1):
        t = ws.cell(r, 1).value
        if not t or str(t).strip().lower() in ('products', 'title'): continue
        if ws.cell(r, 4).fill.fill_type and not any(ws.cell(r, c).value for c in range(4, 18)): continue
        if any(str(ws.cell(r, c).value or '').strip() in ('Installers', 'Programmers', 'Project managers', 'Other agencies') for c in range(1, 18)): break
        t = str(t).strip(); loc = ws.cell(r, 2).value
        locs = EXTRA_LOC.get(t) or ([x.strip() for x in str(loc).split('/') if x.strip()] if loc else [])
        seen, seq = set(), 0
        for c in range(4, 18):
            v = ws.cell(r, c).value
            if not v: continue
            lab = re.sub(r'\s+', ' ', str(v)).strip(); lab = FIX.get(lab, lab); lab = lab[0].upper() + lab[1:]
            if lab.lower() in seen: continue
            f = ws.cell(r, c).fill.fgColor.rgb
            if f not in COL: sys.exit(f'Unknown colour {f} at {ws.cell(r, c).coordinate} ("{lab}"). Decide its owner, add it to COL.')
            seen.add(lab.lower()); seq += 1
            resp, phase = COL[f]
            rows.append(dict(product=t, seq=seq, label=lab, phase=phase, responsibility=resp,
                             gate='none' if phase == 1 else 'after_external',
                             default_agency_type=INTERNAL_TYPE.get(f) or ext_type(t, lab),
                             sheet_colour=f[2:], sheet_cell=ws.cell(r, c).coordinate))
        prod.append(dict(title=t, scope='room' if loc else 'project', location_options='|'.join(locs), sort=len(prod) + 1))
    out = Path(__file__).resolve().parent.parent / 'supabase' / 'seed'
    for name, data in (('products.csv', prod), ('product_steps.csv', rows)):
        with open(out / name, 'w', newline='') as fh:
            w = csv.DictWriter(fh, fieldnames=list(data[0])); w.writeheader(); w.writerows(data)
    print(f'{len(prod)} products, {len(rows)} steps → {out}')

if __name__ == '__main__':
    main(sys.argv[1])
